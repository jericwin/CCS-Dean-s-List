"""Ranking exports (Excel and PDF)."""
import os
import tempfile
import uuid
from io import BytesIO

from flask import current_app

from .domain import STATUS_LABELS, course_program, course_section

COLUMNS = ["Rank", "Full Name", "Program", "Section", "Academic Year", "Semester", "GWA", "Status"]


def ranking_rows(applications):
    return [
        [
            idx,
            app["full_name"],
            course_program(app["course"]),
            course_section(app["course"]),
            app["academic_year"] or "",
            app["semester"] or "",
            f"{app['gwa']:.2f}" if app["gwa"] is not None else "",
            STATUS_LABELS.get(app["status"], app["status"] or ""),
        ]
        for idx, app in enumerate(applications, start=1)
    ]


def to_excel(applications):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Dean's List Ranking"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7A1420")
        cell.alignment = Alignment(horizontal="center")
    for row in ranking_rows(applications):
        ws.append(row)
    for column, width in zip("ABCDEFGH", (7, 32, 10, 12, 15, 11, 8, 18)):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def _pdf_from_docx_template(applications):
    """Fill DL.docx (official letterhead) and convert through MS Word. Windows + Word only."""
    from docx import Document
    from docx2pdf import convert
    import pythoncom

    doc = Document(current_app.config["DL_TEMPLATE_PATH"])
    table = doc.tables[0]
    for row in ranking_rows(applications):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)

    workdir = tempfile.mkdtemp(prefix="dl_export_")
    docx_path = os.path.join(workdir, f"{uuid.uuid4().hex}.docx")
    pdf_path = os.path.join(workdir, f"{uuid.uuid4().hex}.pdf")
    doc.save(docx_path)
    pythoncom.CoInitialize()
    try:
        convert(docx_path, pdf_path)
    finally:
        pythoncom.CoUninitialize()
    try:
        with open(pdf_path, "rb") as fh:
            return BytesIO(fh.read())
    finally:
        for path in (docx_path, pdf_path):
            if os.path.exists(path):
                os.remove(path)
        os.rmdir(workdir)


def _pdf_with_reportlab(applications, subtitle):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    maroon = colors.HexColor("#7A1420")
    orange = colors.HexColor("#E8832A")
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title="Dean's List Ranking")
    story = []
    logo = os.path.join(current_app.static_folder, "logo.png")
    if os.path.exists(logo):
        story.append(Image(logo, width=24 * mm, height=24 * mm))
    story.append(Paragraph("College of Computer Studies — Student Body Organization",
                           ParagraphStyle("org", fontName="Helvetica", fontSize=10, alignment=1, textColor=colors.grey)))
    story.append(Paragraph("Dean's List Ranking",
                           ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=18, alignment=1,
                                          textColor=maroon, spaceBefore=4, spaceAfter=2)))
    story.append(Paragraph(subtitle, ParagraphStyle("sub", fontName="Helvetica", fontSize=10, alignment=1)))
    story.append(Spacer(1, 6 * mm))

    data = [COLUMNS] + ranking_rows(applications)
    table = Table(data, repeatRows=1, colWidths=[14 * mm, 70 * mm, 22 * mm, 26 * mm, 32 * mm, 24 * mm, 18 * mm, 40 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), maroon),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (6, 0), (6, -1), "CENTER"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F4F4")]),
        ("LINEBELOW", (0, 0), (-1, 0), 1.2, orange),
        ("GRID", (0, 1), (-1, -1), 0.25, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    doc.build(story)
    out.seek(0)
    return out


def to_pdf(applications, subtitle):
    engine = current_app.config.get("PDF_EXPORT_ENGINE", "auto")
    if engine in ("auto", "docx") and os.path.exists(current_app.config["DL_TEMPLATE_PATH"]):
        try:
            return _pdf_from_docx_template(applications)
        except Exception as err:  # Word missing, COM failure, non-Windows host...
            if engine == "docx":
                raise
            current_app.logger.info("DL.docx conversion unavailable (%s); using ReportLab.", err)
    return _pdf_with_reportlab(applications, subtitle)
